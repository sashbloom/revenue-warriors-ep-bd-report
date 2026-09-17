"""
Revenue Warriors Dispatch Agent
================================
Reads an uploaded Revenue Warriors workbook, extracts EP/BD sheets,
builds personalized HTML emails, and creates/sends drafts via Graph API.
"""

import json
import logging
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

log = logging.getLogger("revenue_warriors")

CONFIG_PATH = Path(__file__).parent / "config.json"


def load_recipients() -> list[dict]:
    with open(CONFIG_PATH) as f:
        return json.load(f)["recipients"]


def load_cc_emails() -> list[str]:
    """Addresses CC'd on every dispatch (Rahul, Venkat)."""
    with open(CONFIG_PATH) as f:
        return json.load(f).get("cc_emails", [])


def merge_cc(extra: list[str] | None) -> list[str]:
    """Default CC first, then any per-request extras, de-duplicated."""
    merged: list[str] = []
    seen: set[str] = set()
    for e in load_cc_emails() + (extra or []):
        e = e.strip()
        if e and e.lower() not in seen:
            seen.add(e.lower())
            merged.append(e)
    return merged


# ── Workbook parsing ──────────────────────────────────────

def detect_week(wb) -> str:
    """Auto-detect week label from the workbook."""
    weekly = [s for s in wb.sheetnames if s.startswith("Revenue_FY27")]
    if weekly:
        return "Week of " + weekly[0].replace("Revenue_FY27 ", "")
    return "Current Week"


def extract_sheet(wb, sheet_name: str) -> dict:
    """Extract headers + rows from one sheet. Returns {headers, rows, title}."""
    if sheet_name not in wb.sheetnames:
        return {"headers": [], "rows": [], "title": sheet_name}

    ws = wb[sheet_name]
    all_rows = [list(row) for row in ws.iter_rows(values_only=True)]
    if not all_rows:
        return {"headers": [], "rows": [], "title": sheet_name}

    title = all_rows[0][0] or sheet_name

    # Find header row
    h_idx = None
    for i, row in enumerate(all_rows):
        if sum(1 for c in row if c is not None) >= 3:
            h_idx = i
            break
    if h_idx is None:
        return {"headers": [], "rows": [], "title": title}

    headers = all_rows[h_idx]
    rows = all_rows[h_idx + 1 :]

    # Remove "Row Type" column
    tc = next((j for j, h in enumerate(headers) if h and str(h).strip().lower() == "row type"), None)
    if tc is not None:
        headers = headers[:tc] + headers[tc + 1 :]
        rows = [r[:tc] + r[tc + 1 :] for r in rows]

    # Filter empties, trim trailing columns
    rows = [r for r in rows if any(c not in (None, "", 0) for c in r)]
    mx = max(
        (next((j for j in range(len(r) - 1, -1, -1) if r[j] not in (None, "")), -1) + 1)
        for r in [headers] + rows
    ) if rows else len(headers)
    headers = headers[:mx]
    rows = [r[:mx] for r in rows]

    return {"headers": headers, "rows": rows, "title": title}


def extract_all_sheets(file_bytes: bytes) -> dict:
    """Parse the workbook from raw bytes. Returns {week, sheets: {sheet_name: data}}."""
    wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    week = detect_week(wb)

    sheets = {}
    recipients = load_recipients()
    for r in recipients:
        sheet_name = r["sheet"]
        if sheet_name in wb.sheetnames:
            sheets[sheet_name] = extract_sheet(wb, sheet_name)

    wb.close()
    return {"week": week, "sheets": sheets, "total_sheets": len(wb.sheetnames)}


# ── Email building ────────────────────────────────────────

def _fmt(val):
    if val is None or val == "":
        return "—"
    if isinstance(val, (int, float)):
        if val == 0:
            return "—"
        if abs(val) >= 100000:
            return f"₹{val / 100000:,.2f}L"
        if abs(val) >= 1:
            return f"₹{val:,.0f}"
        return f"₹{val:,.2f}"
    return str(val)


def sheet_to_html(headers, rows, title=""):
    h = '<table style="border-collapse:collapse;width:100%;font-family:Calibri,Arial,sans-serif;font-size:13px;margin:16px 0">'
    if title:
        h += f'<caption style="text-align:left;font-weight:bold;font-size:15px;padding:12px 0 8px;color:#1a2942">{title}</caption>'
    h += "<thead><tr>"
    for c in headers:
        h += f'<th style="background:#1a2942;color:#fff;padding:8px 12px;text-align:left;font-size:12px;font-weight:600;border:1px solid #24364f">{c or ""}</th>'
    h += "</tr></thead><tbody>"
    for i, row in enumerate(rows):
        bg = "#f8fafc" if i % 2 == 0 else "#fff"
        h += "<tr>"
        for j, cell in enumerate(row):
            is_num = j > 0 and isinstance(cell, (int, float))
            h += f'<td style="padding:6px 12px;border:1px solid #e2e8f0;background:{bg};text-align:{"right" if is_num else "left"}">{_fmt(cell) if is_num else (cell or "")}</td>'
        h += "</tr>"
    h += "</tbody></table>"
    return h


def build_email(recipient: dict, sheet_data: dict, week: str, cc: list[str]) -> dict:
    """Build one email payload for a recipient."""
    name = recipient["name"]
    rtype = recipient["type"]
    type_label = "Engagement Partner" if rtype == "EP" else "BD Owner"
    counter = "BD owner" if rtype == "EP" else "EP"

    table_html = sheet_to_html(
        sheet_data["headers"], sheet_data["rows"], title=sheet_data["title"]
    )

    subject = f"Revenue Warriors — Your {type_label} Report ({week})"
    body = f"""<p>Hi {name},</p>
<p>Please find your Revenue Warriors <strong>{type_label}</strong> report for <strong>{week}</strong> below.</p>
<p>Your sheet shows committed revenue and upsells broken down by {counter}. Please review and flag any discrepancies by <strong>EOD Wednesday</strong>.</p>
{table_html}
<p style="margin-top:20px;font-size:13px;color:#666">This email was auto-generated from the Revenue Warriors workbook.</p>
<p>Best regards,<br>Isabelle<br>Finance Team, Practus</p>"""

    return {
        "to": [recipient["email"]],
        "cc": cc,
        "subject": subject,
        "body_html": body,
        "name": name,
        "type": rtype,
        "sheet": recipient["sheet"],
        "row_count": len(sheet_data["rows"]),
    }


def build_all_emails(file_bytes: bytes, week_override: str = "", cc: list[str] | None = None) -> dict:
    """Parse workbook + build all email payloads. Returns {week, emails: [...]}."""
    parsed = extract_all_sheets(file_bytes)
    week = week_override or parsed["week"]
    recipients = load_recipients()
    cc = merge_cc(cc)

    emails = []
    for r in recipients:
        sheet_data = parsed["sheets"].get(r["sheet"])
        if sheet_data and sheet_data["headers"]:
            emails.append(build_email(r, sheet_data, week, cc))

    return {"week": week, "emails": emails, "total": len(emails)}
